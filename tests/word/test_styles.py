"""Tests for basic_open_agent_tools.word.styles module."""

from pathlib import Path
from unittest.mock import Mock, patch

import pytest

from basic_open_agent_tools.word.styles import (
    add_page_break,
    apply_bold_to_paragraph,
    apply_heading_style,
    set_paragraph_alignment,
)


class TestApplyHeadingStyle:
    """Test cases for apply_heading_style function."""

    def test_apply_heading_style_basic(self, tmp_path: Path) -> None:
        """Test applying heading style to paragraph."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", True):
            mock_para = Mock()
            mock_para.style = None
            mock_doc = Mock()
            mock_doc.paragraphs = [mock_para]
            mock_doc.save = Mock()

            with patch(
                "basic_open_agent_tools.word.styles.Document",
                return_value=mock_doc,
                create=True,
            ):
                result = apply_heading_style(
                    str(docx_file),
                    paragraph_index=0,
                    heading_level=1,
                    skip_confirm=True,
                )

            # Verify style was set
            assert mock_para.style == "Heading 1"
            mock_doc.save.assert_called_once()

        assert "Applied Heading 1 style" in result

    def test_apply_heading_style_level_2(self, tmp_path: Path) -> None:
        """Test applying heading level 2 style."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", True):
            mock_para = Mock()
            mock_para.style = None
            mock_doc = Mock()
            mock_doc.paragraphs = [mock_para]
            mock_doc.save = Mock()

            with patch(
                "basic_open_agent_tools.word.styles.Document",
                return_value=mock_doc,
                create=True,
            ):
                result = apply_heading_style(
                    str(docx_file),
                    paragraph_index=0,
                    heading_level=2,
                    skip_confirm=True,
                )

            assert mock_para.style == "Heading 2"

        assert "Applied Heading 2 style" in result

    def test_apply_heading_style_invalid_level(self, tmp_path: Path) -> None:
        """Test error with invalid heading level."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", True):
            with pytest.raises(
                ValueError, match="heading_level must be between 1 and 9"
            ):
                apply_heading_style(
                    str(docx_file),
                    paragraph_index=0,
                    heading_level=10,
                    skip_confirm=True,
                )

    def test_apply_heading_style_invalid_paragraph_index(self, tmp_path: Path) -> None:
        """Test error with invalid paragraph index."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", True):
            mock_doc = Mock()
            mock_doc.paragraphs = []

            with patch(
                "basic_open_agent_tools.word.styles.Document",
                return_value=mock_doc,
                create=True,
            ):
                with pytest.raises(IndexError, match="Paragraph index 0 out of range"):
                    apply_heading_style(
                        str(docx_file),
                        paragraph_index=0,
                        heading_level=1,
                        skip_confirm=True,
                    )

    def test_apply_heading_style_file_not_found(self) -> None:
        """Test error when document doesn't exist."""
        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", True):
            with pytest.raises(FileNotFoundError, match="Word document not found"):
                apply_heading_style(
                    "/nonexistent/file.docx",
                    paragraph_index=0,
                    heading_level=1,
                    skip_confirm=True,
                )

    def test_apply_heading_style_python_docx_not_available(self) -> None:
        """Test error when python-docx is not installed."""
        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", False):
            with pytest.raises(ImportError, match="python-docx is required"):
                apply_heading_style("test.docx", 0, 1, skip_confirm=True)


class TestApplyBoldToParagraph:
    """Test cases for apply_bold_to_paragraph function."""

    def test_apply_bold_basic(self, tmp_path: Path) -> None:
        """Test applying bold to paragraph."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", True):
            mock_run = Mock()
            mock_run.bold = False
            mock_para = Mock()
            mock_para.runs = [mock_run]
            mock_doc = Mock()
            mock_doc.paragraphs = [mock_para]
            mock_doc.save = Mock()

            with patch(
                "basic_open_agent_tools.word.styles.Document",
                return_value=mock_doc,
                create=True,
            ):
                result = apply_bold_to_paragraph(
                    str(docx_file), paragraph_index=0, skip_confirm=True
                )

            # Verify bold was applied
            assert mock_run.bold is True
            mock_doc.save.assert_called_once()

        assert "Applied bold formatting" in result

    def test_apply_bold_multiple_runs(self, tmp_path: Path) -> None:
        """Test applying bold to paragraph with multiple runs."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", True):
            mock_run1 = Mock()
            mock_run1.bold = False
            mock_run2 = Mock()
            mock_run2.bold = False
            mock_run3 = Mock()
            mock_run3.bold = False
            mock_para = Mock()
            mock_para.runs = [mock_run1, mock_run2, mock_run3]
            mock_doc = Mock()
            mock_doc.paragraphs = [mock_para]
            mock_doc.save = Mock()

            with patch(
                "basic_open_agent_tools.word.styles.Document",
                return_value=mock_doc,
                create=True,
            ):
                result = apply_bold_to_paragraph(
                    str(docx_file), paragraph_index=0, skip_confirm=True
                )

            # Verify bold was applied to all runs
            assert all(run.bold is True for run in [mock_run1, mock_run2, mock_run3])

        assert "Applied bold formatting" in result

    def test_apply_bold_invalid_paragraph_index(self, tmp_path: Path) -> None:
        """Test error with invalid paragraph index."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", True):
            mock_doc = Mock()
            mock_doc.paragraphs = []

            with patch(
                "basic_open_agent_tools.word.styles.Document",
                return_value=mock_doc,
                create=True,
            ):
                with pytest.raises(IndexError, match="Paragraph index 0 out of range"):
                    apply_bold_to_paragraph(
                        str(docx_file), paragraph_index=0, skip_confirm=True
                    )

    def test_apply_bold_file_not_found(self) -> None:
        """Test error when document doesn't exist."""
        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", True):
            with pytest.raises(FileNotFoundError, match="Word document not found"):
                apply_bold_to_paragraph(
                    "/nonexistent/file.docx", paragraph_index=0, skip_confirm=True
                )

    def test_apply_bold_python_docx_not_available(self) -> None:
        """Test error when python-docx is not installed."""
        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", False):
            with pytest.raises(ImportError, match="python-docx is required"):
                apply_bold_to_paragraph("test.docx", 0, skip_confirm=True)


class TestSetParagraphAlignment:
    """Test cases for set_paragraph_alignment function."""

    def test_set_alignment_left(self, tmp_path: Path) -> None:
        """Test setting paragraph alignment to left."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", True):
            # Mock WD_ALIGN_PARAGRAPH enum
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            mock_para = Mock()
            mock_para.alignment = None
            mock_doc = Mock()
            mock_doc.paragraphs = [mock_para]
            mock_doc.save = Mock()

            with patch(
                "basic_open_agent_tools.word.styles.Document",
                return_value=mock_doc,
                create=True,
            ):
                with patch(
                    "basic_open_agent_tools.word.styles.WD_ALIGN_PARAGRAPH",
                    WD_ALIGN_PARAGRAPH,
                    create=True,
                ):
                    result = set_paragraph_alignment(
                        str(docx_file),
                        paragraph_index=0,
                        alignment="left",
                        skip_confirm=True,
                    )

            # Verify alignment was set
            assert mock_para.alignment is not None
            mock_doc.save.assert_called_once()

        assert "Set alignment" in result

    def test_set_alignment_center(self, tmp_path: Path) -> None:
        """Test setting paragraph alignment to center."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", True):
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            mock_para = Mock()
            mock_para.alignment = None
            mock_doc = Mock()
            mock_doc.paragraphs = [mock_para]
            mock_doc.save = Mock()

            with patch(
                "basic_open_agent_tools.word.styles.Document",
                return_value=mock_doc,
                create=True,
            ):
                with patch(
                    "basic_open_agent_tools.word.styles.WD_ALIGN_PARAGRAPH",
                    WD_ALIGN_PARAGRAPH,
                    create=True,
                ):
                    result = set_paragraph_alignment(
                        str(docx_file),
                        paragraph_index=0,
                        alignment="center",
                        skip_confirm=True,
                    )

            assert mock_para.alignment is not None

        assert "Set alignment" in result

    def test_set_alignment_invalid_value(self, tmp_path: Path) -> None:
        """Test error with invalid alignment value."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", True):
            mock_doc = Mock()
            mock_para = Mock()
            mock_doc.paragraphs = [mock_para]

            with patch(
                "basic_open_agent_tools.word.styles.Document",
                return_value=mock_doc,
                create=True,
            ):
                with pytest.raises(ValueError, match="alignment must be one of"):
                    set_paragraph_alignment(
                        str(docx_file),
                        paragraph_index=0,
                        alignment="invalid",
                        skip_confirm=True,
                    )

    def test_set_alignment_invalid_paragraph_index(self, tmp_path: Path) -> None:
        """Test error with invalid paragraph index."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", True):
            mock_doc = Mock()
            mock_doc.paragraphs = []

            with patch(
                "basic_open_agent_tools.word.styles.Document",
                return_value=mock_doc,
                create=True,
            ):
                with pytest.raises(IndexError, match="Paragraph index 0 out of range"):
                    set_paragraph_alignment(
                        str(docx_file),
                        paragraph_index=0,
                        alignment="left",
                        skip_confirm=True,
                    )

    def test_set_alignment_file_not_found(self) -> None:
        """Test error when document doesn't exist."""
        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", True):
            with pytest.raises(FileNotFoundError, match="Word document not found"):
                set_paragraph_alignment(
                    "/nonexistent/file.docx", 0, "left", skip_confirm=True
                )

    def test_set_alignment_python_docx_not_available(self) -> None:
        """Test error when python-docx is not installed."""
        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", False):
            with pytest.raises(ImportError, match="python-docx is required"):
                set_paragraph_alignment("test.docx", 0, "left", skip_confirm=True)


class TestAddPageBreak:
    """Test cases for add_page_break function."""

    def test_add_page_break_basic(self, tmp_path: Path) -> None:
        """Test adding page break after paragraph."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", True):
            mock_run = Mock()
            mock_run.add_break = Mock()
            mock_para = Mock()
            mock_para.add_run = Mock(return_value=mock_run)
            mock_doc = Mock()
            mock_doc.paragraphs = [mock_para]
            mock_doc.save = Mock()

            with patch(
                "basic_open_agent_tools.word.styles.Document",
                return_value=mock_doc,
                create=True,
            ):
                # Mock BreakType enum
                from docx.enum.text import WD_BREAK

                with patch(
                    "basic_open_agent_tools.word.styles.WD_BREAK",
                    WD_BREAK,
                    create=True,
                ):
                    result = add_page_break(
                        str(docx_file), after_paragraph=0, skip_confirm=True
                    )

            # Verify page break was added
            mock_para.add_run.assert_called_once()
            mock_run.add_break.assert_called_once()
            mock_doc.save.assert_called_once()

        assert "Added page break" in result

    def test_add_page_break_invalid_paragraph_index(self, tmp_path: Path) -> None:
        """Test error with invalid paragraph index."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", True):
            mock_doc = Mock()
            mock_doc.paragraphs = []

            with patch(
                "basic_open_agent_tools.word.styles.Document",
                return_value=mock_doc,
                create=True,
            ):
                with pytest.raises(IndexError, match="Paragraph index 0 out of range"):
                    add_page_break(str(docx_file), after_paragraph=0, skip_confirm=True)

    def test_add_page_break_negative_index(self, tmp_path: Path) -> None:
        """Test error with negative paragraph index."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", True):
            with pytest.raises(
                ValueError, match="after_paragraph must be non-negative"
            ):
                add_page_break(str(docx_file), after_paragraph=-1, skip_confirm=True)

    def test_add_page_break_file_not_found(self) -> None:
        """Test error when document doesn't exist."""
        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", True):
            with pytest.raises(FileNotFoundError, match="Word document not found"):
                add_page_break("/nonexistent/file.docx", 0, skip_confirm=True)

    def test_add_page_break_python_docx_not_available(self) -> None:
        """Test error when python-docx is not installed."""
        with patch("basic_open_agent_tools.word.styles.HAS_PYTHON_DOCX", False):
            with pytest.raises(ImportError, match="python-docx is required"):
                add_page_break("test.docx", 0, skip_confirm=True)
