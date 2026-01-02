"""Tests for basic_open_agent_tools.word.reading module."""

from pathlib import Path
from unittest.mock import MagicMock, Mock, patch

import pytest

from basic_open_agent_tools.word.reading import (
    extract_text_from_docx,
    get_docx_info,
    get_docx_metadata,
    get_docx_paragraphs,
    get_docx_tables,
    search_docx_text,
)


class TestExtractTextFromDocx:
    """Test cases for extract_text_from_docx function."""

    def test_extract_text_basic(self, tmp_path: Path) -> None:
        """Test extracting text from a basic Word document."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            from docx import Document

            # Mock Document
            mock_doc = Mock()
            mock_para1 = Mock()
            mock_para1.text = "First paragraph"
            mock_para2 = Mock()
            mock_para2.text = "Second paragraph"
            mock_para3 = Mock()
            mock_para3.text = ""  # Empty paragraph
            mock_doc.paragraphs = [mock_para1, mock_para2, mock_para3]

            with patch(
                "basic_open_agent_tools.word.reading.Document",
                return_value=mock_doc,
                create=True,
            ):
                result = extract_text_from_docx(str(docx_file))

        assert result == "First paragraph\nSecond paragraph"

    def test_extract_text_unicode_content(self, tmp_path: Path) -> None:
        """Test extracting Unicode text from Word document."""
        docx_file = tmp_path / "unicode.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            mock_doc = Mock()
            mock_para = Mock()
            mock_para.text = "北京 東京 Москва 🚀"
            mock_doc.paragraphs = [mock_para]

            with patch(
                "basic_open_agent_tools.word.reading.Document",
                return_value=mock_doc,
                create=True,
            ):
                result = extract_text_from_docx(str(docx_file))

        assert result == "北京 東京 Москва 🚀"

    def test_extract_text_empty_document(self, tmp_path: Path) -> None:
        """Test extracting text from empty Word document."""
        docx_file = tmp_path / "empty.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            mock_doc = Mock()
            mock_doc.paragraphs = []

            with patch(
                "basic_open_agent_tools.word.reading.Document",
                return_value=mock_doc,
                create=True,
            ):
                result = extract_text_from_docx(str(docx_file))

        assert result == ""

    def test_extract_text_python_docx_not_available(self) -> None:
        """Test error when python-docx is not installed."""
        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", False):
            with pytest.raises(ImportError, match="python-docx is required"):
                extract_text_from_docx("test.docx")

    def test_extract_text_file_not_found(self) -> None:
        """Test error when Word file doesn't exist."""
        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            with pytest.raises(FileNotFoundError, match="Word document not found"):
                extract_text_from_docx("/nonexistent/file.docx")

    def test_extract_text_invalid_file_path_type(self) -> None:
        """Test error when file_path is not a string."""
        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            with pytest.raises(TypeError, match="file_path must be a string"):
                extract_text_from_docx(123)  # type: ignore[arg-type]

    def test_extract_text_empty_file_path(self) -> None:
        """Test error when file_path is empty string."""
        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            with pytest.raises(ValueError, match="file_path cannot be empty"):
                extract_text_from_docx("")

    def test_extract_text_file_too_large(self, tmp_path: Path) -> None:
        """Test error when file is too large."""
        docx_file = tmp_path / "large.docx"
        docx_file.touch()

        # Mock os.path.getsize to return a large size
        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            with patch("os.path.getsize", return_value=60 * 1024 * 1024):  # 60MB
                with pytest.raises(ValueError, match="Word document too large"):
                    extract_text_from_docx(str(docx_file))

    def test_extract_text_permission_error(self, tmp_path: Path) -> None:
        """Test error when file cannot be read."""
        docx_file = tmp_path / "no_permission.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            with patch("os.access", return_value=False):
                with pytest.raises(PermissionError, match="Cannot read Word document"):
                    extract_text_from_docx(str(docx_file))

    def test_extract_text_corrupted_document(self, tmp_path: Path) -> None:
        """Test error when document is corrupted or invalid."""
        docx_file = tmp_path / "corrupted.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            with patch(
                "basic_open_agent_tools.word.reading.Document",
                side_effect=Exception("Invalid document"),
                create=True,
            ):
                with pytest.raises(ValueError, match="Failed to read Word document"):
                    extract_text_from_docx(str(docx_file))


class TestGetDocxParagraphs:
    """Test cases for get_docx_paragraphs function."""

    def test_get_paragraphs_basic(self, tmp_path: Path) -> None:
        """Test getting paragraphs from Word document."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            mock_doc = Mock()
            mock_para1 = Mock()
            mock_para1.text = "Paragraph one"
            mock_para2 = Mock()
            mock_para2.text = "Paragraph two"
            mock_para3 = Mock()
            mock_para3.text = "Paragraph three"
            mock_doc.paragraphs = [mock_para1, mock_para2, mock_para3]

            with patch(
                "basic_open_agent_tools.word.reading.Document",
                return_value=mock_doc,
                create=True,
            ):
                result = get_docx_paragraphs(str(docx_file))

        assert result == ["Paragraph one", "Paragraph two", "Paragraph three"]

    def test_get_paragraphs_with_empty_paragraphs(self, tmp_path: Path) -> None:
        """Test getting paragraphs including empty ones."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            mock_doc = Mock()
            mock_para1 = Mock()
            mock_para1.text = "First"
            mock_para2 = Mock()
            mock_para2.text = ""  # Empty paragraph
            mock_para3 = Mock()
            mock_para3.text = "Third"
            mock_doc.paragraphs = [mock_para1, mock_para2, mock_para3]

            with patch(
                "basic_open_agent_tools.word.reading.Document",
                return_value=mock_doc,
                create=True,
            ):
                result = get_docx_paragraphs(str(docx_file))

        assert result == ["First", "", "Third"]

    def test_get_paragraphs_empty_document(self, tmp_path: Path) -> None:
        """Test getting paragraphs from empty document."""
        docx_file = tmp_path / "empty.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            mock_doc = Mock()
            mock_doc.paragraphs = []

            with patch(
                "basic_open_agent_tools.word.reading.Document",
                return_value=mock_doc,
                create=True,
            ):
                result = get_docx_paragraphs(str(docx_file))

        assert result == []

    def test_get_paragraphs_python_docx_not_available(self) -> None:
        """Test error when python-docx is not installed."""
        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", False):
            with pytest.raises(ImportError, match="python-docx is required"):
                get_docx_paragraphs("test.docx")


class TestGetDocxTables:
    """Test cases for get_docx_tables function."""

    def test_get_tables_basic(self, tmp_path: Path) -> None:
        """Test getting tables from Word document."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            # Mock table with 2x2 cells
            mock_cell_1_1 = Mock()
            mock_cell_1_1.text = "Header 1"
            mock_cell_1_2 = Mock()
            mock_cell_1_2.text = "Header 2"
            mock_cell_2_1 = Mock()
            mock_cell_2_1.text = "Data 1"
            mock_cell_2_2 = Mock()
            mock_cell_2_2.text = "Data 2"

            mock_row_1 = Mock()
            mock_row_1.cells = [mock_cell_1_1, mock_cell_1_2]
            mock_row_2 = Mock()
            mock_row_2.cells = [mock_cell_2_1, mock_cell_2_2]

            mock_table = Mock()
            mock_table.rows = [mock_row_1, mock_row_2]

            mock_doc = Mock()
            mock_doc.tables = [mock_table]

            with patch(
                "basic_open_agent_tools.word.reading.Document",
                return_value=mock_doc,
                create=True,
            ):
                result = get_docx_tables(str(docx_file))

        expected = [[["Header 1", "Header 2"], ["Data 1", "Data 2"]]]
        assert result == expected

    def test_get_tables_multiple_tables(self, tmp_path: Path) -> None:
        """Test getting multiple tables from document."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            # First table
            mock_cell_1 = Mock()
            mock_cell_1.text = "Table 1"
            mock_row_1 = Mock()
            mock_row_1.cells = [mock_cell_1]
            mock_table_1 = Mock()
            mock_table_1.rows = [mock_row_1]

            # Second table
            mock_cell_2 = Mock()
            mock_cell_2.text = "Table 2"
            mock_row_2 = Mock()
            mock_row_2.cells = [mock_cell_2]
            mock_table_2 = Mock()
            mock_table_2.rows = [mock_row_2]

            mock_doc = Mock()
            mock_doc.tables = [mock_table_1, mock_table_2]

            with patch(
                "basic_open_agent_tools.word.reading.Document",
                return_value=mock_doc,
                create=True,
            ):
                result = get_docx_tables(str(docx_file))

        assert len(result) == 2
        assert result[0] == [["Table 1"]]
        assert result[1] == [["Table 2"]]

    def test_get_tables_no_tables(self, tmp_path: Path) -> None:
        """Test getting tables when document has none."""
        docx_file = tmp_path / "no_tables.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            mock_doc = Mock()
            mock_doc.tables = []

            with patch(
                "basic_open_agent_tools.word.reading.Document",
                return_value=mock_doc,
                create=True,
            ):
                result = get_docx_tables(str(docx_file))

        assert result == []

    def test_get_tables_python_docx_not_available(self) -> None:
        """Test error when python-docx is not installed."""
        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", False):
            with pytest.raises(ImportError, match="python-docx is required"):
                get_docx_tables("test.docx")


class TestSearchDocxText:
    """Test cases for search_docx_text function."""

    def test_search_text_found(self, tmp_path: Path) -> None:
        """Test searching for text in Word document."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            mock_doc = Mock()
            mock_para1 = Mock()
            mock_para1.text = "This is a test paragraph"
            mock_para2 = Mock()
            mock_para2.text = "Another test sentence"
            mock_para3 = Mock()
            mock_para3.text = "No match here"
            mock_doc.paragraphs = [mock_para1, mock_para2, mock_para3]

            with patch(
                "basic_open_agent_tools.word.reading.Document",
                return_value=mock_doc,
                create=True,
            ):
                result = search_docx_text(str(docx_file), "test", case_sensitive=False)

        assert len(result) == 2
        assert result[0]["paragraph_text"] == "This is a test paragraph"
        assert result[1]["paragraph_text"] == "Another test sentence"

    def test_search_text_case_sensitive(self, tmp_path: Path) -> None:
        """Test case-sensitive search."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            mock_doc = Mock()
            mock_para1 = Mock()
            mock_para1.text = "Test with capital T"
            mock_para2 = Mock()
            mock_para2.text = "test with lowercase t"
            mock_doc.paragraphs = [mock_para1, mock_para2]

            with patch(
                "basic_open_agent_tools.word.reading.Document",
                return_value=mock_doc,
                create=True,
            ):
                result = search_docx_text(str(docx_file), "Test", case_sensitive=True)

        assert len(result) == 1
        assert result[0]["paragraph_text"] == "Test with capital T"

    def test_search_text_case_insensitive(self, tmp_path: Path) -> None:
        """Test case-insensitive search."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            mock_doc = Mock()
            mock_para1 = Mock()
            mock_para1.text = "Test with capital T"
            mock_para2 = Mock()
            mock_para2.text = "test with lowercase t"
            mock_doc.paragraphs = [mock_para1, mock_para2]

            with patch(
                "basic_open_agent_tools.word.reading.Document",
                return_value=mock_doc,
                create=True,
            ):
                result = search_docx_text(str(docx_file), "test", case_sensitive=False)

        assert len(result) == 2
        assert result[0]["paragraph_text"] == "Test with capital T"
        assert result[1]["paragraph_text"] == "test with lowercase t"

    def test_search_text_not_found(self, tmp_path: Path) -> None:
        """Test searching for text that doesn't exist."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            mock_doc = Mock()
            mock_para = Mock()
            mock_para.text = "This has no match"
            mock_doc.paragraphs = [mock_para]

            with patch(
                "basic_open_agent_tools.word.reading.Document",
                return_value=mock_doc,
                create=True,
            ):
                result = search_docx_text(
                    str(docx_file), "nonexistent", case_sensitive=False
                )

        assert result == []

    def test_search_text_python_docx_not_available(self) -> None:
        """Test error when python-docx is not installed."""
        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", False):
            with pytest.raises(ImportError, match="python-docx is required"):
                search_docx_text("test.docx", "search", case_sensitive=False)


class TestGetDocxMetadata:
    """Test cases for get_docx_metadata function."""

    def test_get_metadata_basic(self, tmp_path: Path) -> None:
        """Test getting metadata from Word document."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            from datetime import datetime

            mock_core_props = Mock()
            mock_core_props.title = "Test Document"
            mock_core_props.author = "Test Author"
            mock_core_props.subject = "Test Subject"
            mock_core_props.keywords = "test, keywords"
            mock_core_props.created = datetime(2024, 1, 1)
            mock_core_props.modified = datetime(2024, 1, 2)

            mock_doc = Mock()
            mock_doc.core_properties = mock_core_props

            with patch(
                "basic_open_agent_tools.word.reading.Document",
                return_value=mock_doc,
                create=True,
            ):
                result = get_docx_metadata(str(docx_file))

        assert result["title"] == "Test Document"
        assert result["author"] == "Test Author"
        assert result["subject"] == "Test Subject"
        assert result["keywords"] == "test, keywords"
        assert result["created"] == "2024-01-01 00:00:00"
        assert result["modified"] == "2024-01-02 00:00:00"

    def test_get_metadata_empty_fields(self, tmp_path: Path) -> None:
        """Test getting metadata with empty fields."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            mock_core_props = Mock()
            mock_core_props.title = None
            mock_core_props.author = None
            mock_core_props.subject = None
            mock_core_props.keywords = None
            mock_core_props.created = None
            mock_core_props.modified = None

            mock_doc = Mock()
            mock_doc.core_properties = mock_core_props

            with patch(
                "basic_open_agent_tools.word.reading.Document",
                return_value=mock_doc,
                create=True,
            ):
                result = get_docx_metadata(str(docx_file))

        assert result["title"] == ""
        assert result["author"] == ""
        assert result["subject"] == ""
        assert result["keywords"] == ""
        assert result["created"] == ""
        assert result["modified"] == ""

    def test_get_metadata_python_docx_not_available(self) -> None:
        """Test error when python-docx is not installed."""
        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", False):
            with pytest.raises(ImportError, match="python-docx is required"):
                get_docx_metadata("test.docx")


class TestGetDocxInfo:
    """Test cases for get_docx_info function."""

    def test_get_info_basic(self, tmp_path: Path) -> None:
        """Test getting general info from Word document."""
        docx_file = tmp_path / "test.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            from datetime import datetime

            mock_core_props = Mock()
            mock_core_props.title = "Test"
            mock_core_props.author = "Author"
            mock_core_props.subject = "Subject"
            mock_core_props.keywords = "keywords"
            mock_core_props.created = datetime(2024, 1, 1)
            mock_core_props.modified = datetime(2024, 1, 2)

            mock_doc = Mock()
            mock_para1 = Mock()
            mock_para1.text = "First paragraph"
            mock_para2 = Mock()
            mock_para2.text = "Second paragraph"
            mock_doc.paragraphs = [mock_para1, mock_para2]

            # Mock table
            mock_table = Mock()
            mock_doc.tables = [mock_table]

            # Mock core properties
            mock_doc.core_properties = mock_core_props

            with patch(
                "basic_open_agent_tools.word.reading.Document",
                return_value=mock_doc,
                create=True,
            ):
                with patch("os.path.getsize", return_value=1024):
                    result = get_docx_info(str(docx_file))

        assert result["paragraph_count"] == 2
        assert result["table_count"] == 1
        assert result["file_size_bytes"] == 1024
        assert "metadata" in result
        assert result["metadata"]["title"] == "Test"

    def test_get_info_empty_document(self, tmp_path: Path) -> None:
        """Test getting info from empty document."""
        docx_file = tmp_path / "empty.docx"
        docx_file.touch()

        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", True):
            mock_core_props = Mock()
            mock_core_props.title = None
            mock_core_props.author = None
            mock_core_props.subject = None
            mock_core_props.keywords = None
            mock_core_props.created = None
            mock_core_props.modified = None

            mock_doc = Mock()
            mock_doc.paragraphs = []
            mock_doc.tables = []
            mock_doc.core_properties = mock_core_props

            with patch(
                "basic_open_agent_tools.word.reading.Document",
                return_value=mock_doc,
                create=True,
            ):
                with patch("os.path.getsize", return_value=512):
                    result = get_docx_info(str(docx_file))

        assert result["paragraph_count"] == 0
        assert result["table_count"] == 0
        assert result["file_size_bytes"] == 512
        assert "metadata" in result

    def test_get_info_python_docx_not_available(self) -> None:
        """Test error when python-docx is not installed."""
        with patch("basic_open_agent_tools.word.reading.HAS_PYTHON_DOCX", False):
            with pytest.raises(ImportError, match="python-docx is required"):
                get_docx_info("test.docx")
